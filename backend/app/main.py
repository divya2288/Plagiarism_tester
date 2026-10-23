import os, re, io, hashlib, secrets
from dotenv import load_dotenv
from datetime import datetime, timedelta, timezone
from pathlib import Path
from difflib import SequenceMatcher
from typing import Optional
from fastapi.middleware.cors import CORSMiddleware

import fitz
from docx import Document as DocxDocument
from fastapi import FastAPI, Depends, File, Form, HTTPException, UploadFile, Header

from fastapi.responses import StreamingResponse
from jose import jwt, JWTError
from pydantic import BaseModel, EmailStr
from sqlalchemy import create_engine, String, Text, Integer, Float, DateTime, ForeignKey, Boolean
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, relationship, sessionmaker, Session
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from xml.sax.saxutils import escape
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from rapidfuzz.fuzz import ratio

load_dotenv(Path(__file__).resolve().parent.parent / ".env")
semantic_model = None

if os.getenv("USE_SENTENCE_TRANSFORMERS", "false").lower() == "true":
    semantic_model = SentenceTransformer("all-MiniLM-L6-v2")
DATABASE_URL = os.getenv("DATABASE_URL")
JWT_SECRET = os.getenv("JWT_SECRET", "dev-only-change-me-before-deploying")
JWT_ALGORITHM = "HS256"
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", "uploads"))
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "10"))
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql+psycopg://", 1)
elif DATABASE_URL.startswith("postgresql://"):
    DATABASE_URL = DATABASE_URL.replace("postgresql://", "postgresql+psycopg://", 1)

engine = create_engine(DATABASE_URL, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)

class Base(DeclarativeBase):
    pass

class User(Base):
    __tablename__ = "users"

    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    email: Mapped[str] = mapped_column(String(255), unique=True, index=True)
    password_hash: Mapped[str] = mapped_column(String(255))
    role: Mapped[str] = mapped_column(String(20), default="user", nullable=False)
    created_at: Mapped[datetime] = mapped_column(
        DateTime, default=datetime.utcnow
    )

class ReferenceDocument(Base):
    __tablename__ = "reference_documents"

    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"), nullable=True)
    name: Mapped[str] = mapped_column(String(255))
    text: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(
        DateTime, default=datetime.utcnow
    )

class Scan(Base):
    __tablename__ = "scans"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    document_name: Mapped[str] = mapped_column(String(255))
    submitted_text: Mapped[str] = mapped_column(Text)
    word_count: Mapped[int] = mapped_column(Integer)
    similarity_score: Mapped[float] = mapped_column(Float, default=0)
    matched_word_count: Mapped[int] = mapped_column(Integer, default=0)
    match_count: Mapped[int] = mapped_column(Integer, default=0)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    matches: Mapped[list["Match"]] = relationship(back_populates="scan", cascade="all, delete-orphan")

class Match(Base):
    __tablename__ = "matches"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    scan_id: Mapped[int] = mapped_column(ForeignKey("scans.id", ondelete="CASCADE"), index=True)
    category: Mapped[str] = mapped_column(String(32))
    submitted_excerpt: Mapped[str] = mapped_column(Text)
    source_name: Mapped[str] = mapped_column(String(255))
    source_excerpt: Mapped[str] = mapped_column(Text)
    similarity: Mapped[float] = mapped_column(Float)
    start_char: Mapped[int] = mapped_column(Integer)
    end_char: Mapped[int] = mapped_column(Integer)
    scan: Mapped[Scan] = relationship(back_populates="matches")

Base.metadata.create_all(bind=engine)

app = FastAPI(title="PSM Plagiarism Detection API", version="1.0.0")
CORS_ORIGINS = [
    origin.strip()
    for origin in os.getenv(
        "CORS_ORIGINS",
        "http://localhost:5173"
    ).split(",")
    if origin.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

def hash_password(password: str, salt: Optional[str] = None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 310_000).hex()
    return f"{salt}${digest}"

def verify_password(password: str, stored: str):
    try:
        salt, expected = stored.split("$", 1)
        return secrets.compare_digest(hash_password(password, salt).split("$", 1)[1], expected)
    except ValueError:
        return False

def make_token(user: User):
    payload = {"sub": str(user.id), "email": user.email, "exp": datetime.now(timezone.utc) + timedelta(hours=12)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

def current_user(authorization: Optional[str] = Header(default=None), db: Session = Depends(get_db)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Sign in to continue.")
    try:
        payload = jwt.decode(authorization[7:], JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user_id = int(payload["sub"])
    except (JWTError, KeyError, ValueError):
        raise HTTPException(401, "Your session is invalid or expired. Please sign in again.")
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(401, "User account not found.")
    return user

class RegisterBody(BaseModel):
    email: EmailStr
    password: str
class LoginBody(BaseModel):
    email: EmailStr
    password: str
class TextScanBody(BaseModel):
    text: str
    document_name: str = "Pasted text.txt"

def extract_text(filename: str, raw: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, "Unsupported file type. Upload PDF, DOCX, or TXT.")
    try:
        if ext == ".pdf":
            with fitz.open(stream=raw, filetype="pdf") as doc:
                text = "\n".join(page.get_text("text") for page in doc)
        elif ext == ".docx":
            doc = DocxDocument(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        else:
            text = raw.decode("utf-8-sig")
    except Exception:
        raise HTTPException(400, "The file could not be read. Check that it is a valid, uncorrupted document.")
    text = text.strip()
    if not text:
        raise HTTPException(400, "No readable text was found. Scanned-image PDFs need OCR, which is not included in this version.")
    return text

def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text))

def sentences_with_offsets(text: str):
    # Keep source offsets so the UI can highlight passages in the submitted document.
    out = []
    for m in re.finditer(r"[^.!?\n]+(?:[.!?]+|$)", text):
        s = m.group().strip()
        if len(s.split()) >= 4:
            start = m.start() + len(m.group()) - len(m.group().lstrip())
            out.append((s, start, start + len(s)))
    return out

def detect_matches(submitted: str, references: list[ReferenceDocument]):
    """Compare submitted sentences with reference documents.

    Exact matches take priority. TF-IDF word/character similarity detects
    likely near-duplicates. Optional Sentence Transformers embeddings detect
    semantic matches when USE_SENTENCE_TRANSFORMERS=true is configured.
    """
    sentences = sentences_with_offsets(submitted)
    if not sentences or not references:
        return []

    norm = lambda s: re.sub(r"\W+", " ", s.lower()).strip()
    ref_sentences = []
    for ref in references:
        for ref_sentence, _, _ in sentences_with_offsets(ref.text):
            ref_sentences.append((ref_sentence, ref.name, ref.id))

    if not ref_sentences:
        return []

    submitted_texts = [norm(sentence) for sentence, _, _ in sentences]
    reference_texts = [norm(sentence) for sentence, _, _ in ref_sentences]
    n_submitted = len(sentences)

    word_similarities = None
    char_similarities = None
    semantic_similarities = None

    try:
        vectorizer = TfidfVectorizer(
            analyzer="word", ngram_range=(1, 2), sublinear_tf=True
        )
        matrix = vectorizer.fit_transform(submitted_texts + reference_texts)
        word_similarities = cosine_similarity(
            matrix[:n_submitted], matrix[n_submitted:]
        )
    except ValueError:
        pass

    try:
        vectorizer = TfidfVectorizer(
            analyzer="char_wb", ngram_range=(3, 5), sublinear_tf=True
        )
        matrix = vectorizer.fit_transform(submitted_texts + reference_texts)
        char_similarities = cosine_similarity(
            matrix[:n_submitted], matrix[n_submitted:]
        )
    except ValueError:
        pass

    if semantic_model is not None:
      try:
          embeddings = semantic_model.encode(
            [sentence for sentence, _, _ in sentences]
            + [item[0] for item in ref_sentences],
            normalize_embeddings=True,
          )

          semantic_similarities = cosine_similarity(
            embeddings[:n_submitted],
            embeddings[n_submitted:]
          )
      except Exception as exc:
        print(f"Semantic matching unavailable: {exc}")

    matches = []
    for i, (sentence, start, end) in enumerate(sentences):
        normalized_sentence = submitted_texts[i]
        if len(normalized_sentence.split()) < 4:
            continue

        exact = next(
            (
                (ref_sentence, name, ref_id)
                for ref_sentence, name, ref_id in ref_sentences
                if norm(ref_sentence) == normalized_sentence
            ),
            None,
        )

        best = exact
        category = "exact" if exact else None
        score = 100.0 if exact else 0.0

        if not exact:
            for j, (ref_sentence, name, ref_id) in enumerate(ref_sentences):
                lexical_scores = [ratio(normalized_sentence, reference_texts[j])]
                if word_similarities is not None:
                    lexical_scores.append(float(word_similarities[i, j]) * 100)
                if char_similarities is not None:
                    lexical_scores.append(float(char_similarities[i, j]) * 100)
                lexical_score = max(lexical_scores)

                if lexical_score >= 75 and lexical_score > score:
                    best = (ref_sentence, name, ref_id)
                    category = "near-duplicate"
                    score = lexical_score

                if semantic_similarities is not None:
                    semantic_score = float(semantic_similarities[i, j]) * 100
                    if semantic_score >= 72 and semantic_score > score:
                        best = (ref_sentence, name, ref_id)
                        category = "semantic"
                        score = semantic_score

        if best:
            ref_sentence, source_name, _ = best
            matches.append({
                "category": category,
                "submitted_excerpt": sentence,
                "source_name": source_name,
                "source_excerpt": ref_sentence,
                "similarity": round(score, 1),
                "start_char": start,
                "end_char": end,
            })

    # Keep the strongest match when submitted passages overlap.
    matches.sort(key=lambda item: (item["start_char"], -item["similarity"]))
    filtered = []
    for match in matches:
        if filtered and match["start_char"] < filtered[-1]["end_char"]:
            if match["similarity"] > filtered[-1]["similarity"]:
                filtered[-1] = match
            continue
        filtered.append(match)

    return filtered

def build_result(scan: Scan):
    return {
        "id": scan.id, "document_name": scan.document_name, "created_at": scan.created_at.isoformat(),
        "word_count": scan.word_count, "similarity_score": scan.similarity_score,
        "matched_word_count": scan.matched_word_count, "match_count": scan.match_count,
        "text": scan.submitted_text,
        "matches": [{"id": m.id, "category": m.category, "submitted_excerpt": m.submitted_excerpt,
                     "source_name": m.source_name, "source_excerpt": m.source_excerpt,
                     "similarity": m.similarity, "start_char": m.start_char, "end_char": m.end_char} for m in scan.matches]
    }


def run_scan(db: Session, user: User, text: str, name: str):
    # Compare the submitted document against the shared reference library
    refs = db.query(ReferenceDocument).all()
    matches = detect_matches(text, refs)

    total_words = max(word_count(text), 1)

    # Count unique words covered by matched passages
    matched_word_positions = set()

    for match in matches:
        start = match["start_char"]
        end = match["end_char"]

        for word_match in re.finditer(r"\b[\w'-]+\b", text[start:end]):
            matched_word_positions.add(
                (
                    start + word_match.start(),
                    start + word_match.end(),
                )
            )

    matched_words = len(matched_word_positions)

    # Calculate similarity using each match's word count and similarity
    weighted_matched_words = 0.0

    for match in matches:
        start = match["start_char"]
        end = match["end_char"]

        passage_word_count = len(
            list(
                re.finditer(
                    r"\b[\w'-]+\b",
                    text[start:end],
                )
            )
        )

        weighted_matched_words += (
            passage_word_count * (match["similarity"] / 100)
        )

    score = round(
        min(100.0, weighted_matched_words / total_words * 100),
        1,
    )

    # Save the scan
    scan = Scan(
        user_id=user.id,
        document_name=Path(name).name[:255],
        submitted_text=text,
        word_count=word_count(text),
        similarity_score=score,
        matched_word_count=matched_words,
        match_count=len(matches),
    )

    db.add(scan)
    db.flush()

    # Save individual matched passages
    for match in matches:
        db.add(Match(scan_id=scan.id, **match))

    db.commit()
    db.refresh(scan)

    return build_result(scan)


@app.get("/api/health")
def health():
    return {"status": "ok", "service": "PSM Plagiarism Detection API"}

@app.post("/api/auth/register")
def register(body: RegisterBody, db: Session = Depends(get_db)):
    if len(body.password) < 8:
        raise HTTPException(400, "Password must contain at least 8 characters.")
    email = body.email.lower()
    if db.query(User).filter_by(email=email).first():
        raise HTTPException(409, "An account with this email already exists.")
    user = User(email=email, password_hash=hash_password(body.password))
    db.add(user); db.commit(); db.refresh(user)
    return {"token": make_token(user), "email": user.email}

@app.post("/api/auth/login")
def login(body: LoginBody, db: Session = Depends(get_db)):
    user = db.query(User).filter_by(email=body.email.lower()).first()
    if not user or not verify_password(body.password, user.password_hash):
        raise HTTPException(401, "Incorrect email or password.")
    return {"token": make_token(user), "email": user.email}

@app.get("/api/auth/me")
def me(user: User = Depends(current_user)):
    return {"id": user.id, "email": user.email}

@app.post("/api/references")
async def add_reference(
    file: UploadFile = File(...),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    name = Path(file.filename or "reference.txt").name

    if Path(name).suffix.lower() not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            400, "Upload a PDF, DOCX, or TXT reference file."
        )

    raw = await file.read(MAX_UPLOAD_BYTES + 1)

    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            413, f"File exceeds the {MAX_UPLOAD_MB} MB limit."
        )

    text = extract_text(name, raw)

    if not text or not text.strip():
        raise HTTPException(
            400, "The uploaded document contains no readable text."
        )

    ref = ReferenceDocument(
        name=name[:255],
        text=text,
    )

    db.add(ref)
    db.commit()
    db.refresh(ref)

    return {
        "id": ref.id,
        "name": ref.name,
        "word_count": word_count(ref.text),
        "message": "Reference document added.",
    }
@app.get("/api/references")
def list_references(
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    refs = (
        db.query(ReferenceDocument)
        .order_by(ReferenceDocument.created_at.desc())
        .all()
    )

    return [
        {
            "id": r.id,
            "name": r.name,
            "word_count": word_count(r.text),
            "created_at": r.created_at.isoformat(),
        }
        for r in refs
    ]
@app.delete("/api/references/{reference_id}")
def delete_reference(
    reference_id: int,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    ref = db.get(ReferenceDocument, reference_id)

    if not ref:
        raise HTTPException(404, "Reference not found.")

    db.delete(ref)
    db.commit()

    return {"message": "Reference deleted."}


@app.post("/api/scans")
async def scan_upload(
    file: UploadFile = File(...),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    name = Path(file.filename or "upload.txt").name

    if Path(name).suffix.lower() not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Upload a PDF, DOCX, or TXT file.",
        )

    raw = await file.read(MAX_UPLOAD_BYTES + 1)

    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File exceeds the {MAX_UPLOAD_MB} MB limit.",
        )

    if not raw.strip():
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    try:
        text = extract_text(name, raw)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to read this file. Please upload a valid PDF, DOCX, or TXT document.",
        )

    if not text or not text.strip():
        raise HTTPException(
            status_code=400,
            detail="No readable text was found in the document.",
        )

    return run_scan(db, user, text, name)


@app.post("/api/scans/text")
def scan_text(body: TextScanBody, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not body.text.strip(): raise HTTPException(400, "Paste some text before scanning.")
    if len(body.text) > 1_000_000: raise HTTPException(413, "Pasted text is too long.")
    return run_scan(db, user, body.text, body.document_name or "Pasted text.txt")

@app.get("/api/scans")
def history(user: User = Depends(current_user), db: Session = Depends(get_db)):
    scans = db.query(Scan).filter_by(user_id=user.id).order_by(Scan.created_at.desc()).all()
    return [{"id": s.id, "document_name": s.document_name, "created_at": s.created_at.isoformat(),
             "word_count": s.word_count, "similarity_score": s.similarity_score,
             "match_count": s.match_count} for s in scans]

@app.get("/api/scans/{scan_id}")
def get_scan(scan_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    scan = db.query(Scan).filter_by(id=scan_id, user_id=user.id).first()
    if not scan: raise HTTPException(404, "Scan report not found.")
    return build_result(scan)

@app.delete("/api/scans/{scan_id}")
def delete_scan(scan_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    scan = db.query(Scan).filter_by(id=scan_id, user_id=user.id).first()
    if not scan: raise HTTPException(404, "Scan report not found.")
    db.delete(scan); db.commit()
    return {"message": "Scan and report deleted."}

@app.get("/api/scans/{scan_id}/pdf")
def download_pdf(scan_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    scan = db.query(Scan).filter_by(id=scan_id, user_id=user.id).first()
    if not scan: raise HTTPException(404, "Scan report not found.")
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = [Paragraph("PSM Plagiarism Detection Report", styles["Title"]), Spacer(1, 12)]
    for label, value in [("Document", scan.document_name), ("Scan date", scan.created_at.strftime("%Y-%m-%d %H:%M UTC")),
                         ("Word count", str(scan.word_count)), ("Similarity score", f"{scan.similarity_score:.1f}%"),
                         ("Matched words", str(scan.matched_word_count)), ("Matching passages", str(scan.match_count))]:
        story.append(Paragraph(f"<b>{escape(label)}:</b> {escape(value)}", styles["BodyText"]))
        story.append(Spacer(1, 5))
    story += [Spacer(1, 12), Paragraph("Matches", styles["Heading2"])]
    if scan.matches:
        rows = [["Type", "Source", "Match (%)", "Submitted passage"]]
        for m in scan.matches:
            rows.append([m.category, Paragraph(escape(m.source_name), styles["BodyText"]),
                         f"{m.similarity:.1f}", Paragraph(escape(m.submitted_excerpt[:500]), styles["BodyText"])])
        table = Table(rows, colWidths=[75, 100, 55, 245], repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#243b53")),
                                   ("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.4,colors.grey),
                                   ("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),6)]))
        story.append(table)
    else:
        story.append(Paragraph("No matching passages were found in the current reference database.", styles["BodyText"]))
    story += [Spacer(1, 16), Paragraph("Interpretation: Similarity is an indicator, not definitive proof of plagiarism. Results depend on the reference collection and detection methods.", styles["Italic"])]
    doc.build(story); buffer.seek(0)
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(scan.document_name).stem)[:50] or "report"
    return StreamingResponse(buffer, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}_plagiarism_report.pdf"'})